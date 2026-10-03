"""Regression tests for defects found in QA review of the prototype (IDs from the QA report)."""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from shortlist.config import Settings
from shortlist.rules import applicant_count, is_stale, passes
from shortlist.scoring import has_word, job_flags, keyword_match
from shortlist.tracker import COL, add_file, init_tracker, mark_applied, new_workbook, put, read_jobs

from .conftest import TODAY, make_job

ROOT = Path(__file__).resolve().parent.parent


def write(tmp_path, jobs):
    p = tmp_path / "jobs.json"
    p.write_text(json.dumps(jobs), encoding="utf-8")
    return p


# ---------------------------------------------------------------- B-001: "1.2K applicants"
@pytest.mark.parametrize(
    ("text", "expected"),
    [("1.2K applicants", 1200), ("1.2k", 1200), ("3K", 3000), ("0.5k applicants", 500), ("1,234", 1234)],
)
def test_b001_k_suffix_counts(text, expected):
    assert applicant_count(text) == expected


def test_b001_k_suffix_is_a_real_count_for_the_rule():
    # 5 days old with 1.2K applicants is crowded (over 500), so it is dropped for that reason
    ok, _, reason = passes({"Date Posted": "2026-09-24", "Applicants": "1.2K applicants"}, TODAY)
    assert not ok and "1200 applicants" in reason
    assert applicant_count("Over 1.2K") is None  # capped counts stay unknown


# ---------------------------------------------------------------- B-003: calendar days, documented
def test_b003_fresh_days_are_calendar_days():
    # posted two calendar days ago (could be up to ~71 hours): still "fresh"
    assert passes({"Date Posted": "2026-09-27", "Applicants": "Over 200"}, TODAY) == (True, 2, "fresh")
    assert passes({"Date Posted": "2026-09-26", "Applicants": "Over 200"}, TODAY)[0] is False


def test_b003_documented_in_config_and_readme():
    assert "calendar days" in (Settings.__doc__ or "")
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    assert "previous 2 calendar days" in readme


# ---------------------------------------------------------------- B-004: formula injection
@pytest.mark.parametrize("text", ['=HYPERLINK("https://evil.example.com","Click")', "+1+1", "-2+3", "@SUM(A1)"])
def test_b004_formula_like_text_written_as_text(tmp_path, text):
    settings = Settings(base_dir=tmp_path)
    init_tracker(settings)
    add_file(settings, write(tmp_path, [make_job("Acme", text)]))
    ws = load_workbook(settings.tracker_path)["Tracker"]
    cell = ws.cell(row=2, column=COL["Role Title"])
    assert cell.data_type == "s"
    assert cell.value == text
    assert cell.quotePrefix


def test_b004_put_helper():
    ws = new_workbook(Settings())["Tracker"]
    assert put(ws, 5, 5, "=1+1").data_type == "s"
    assert put(ws, 5, 6, "plain").quotePrefix is False


# ---------------------------------------------------------------- B-005: only http(s) links become hyperlinks
@pytest.mark.parametrize(
    ("url", "linked"),
    [
        ("https://jobs.example.com/1", True),
        ("javascript:alert(1)", False),
        ("data:text/html,<b>x</b>", False),
        ("file:///C:/Windows/system32/calc.exe", False),
        ("https://", False),
    ],
)
def test_b005_hyperlinks_only_for_http_urls(tmp_path, url, linked):
    settings = Settings(base_dir=tmp_path)
    init_tracker(settings)
    add_file(settings, write(tmp_path, [make_job("Acme", "PM", **{"Job URL": url})]))
    wb = load_workbook(settings.tracker_path)
    cell = wb["Tracker"].cell(row=2, column=COL["Job URL"])
    assert (cell.hyperlink is not None) is linked
    dashboard_links = [c.hyperlink.target for row in wb["Dashboard"].iter_rows() for c in row if c.hyperlink]
    assert all(t.startswith("https://") for t in dashboard_links)


# ---------------------------------------------------------------- future posting dates
def test_future_posted_date_is_dropped():
    assert passes({"Date Posted": "2026-09-30", "Applicants": "10"}, TODAY) == (
        False, -1, "posted date in the future")
    assert is_stale({"Date Posted": "2026-10-05", "Applicants": "10"}, TODAY)


# ---------------------------------------------------------------- junior flag vs Director/Manager/Lead/Head
@pytest.mark.parametrize(
    ("title", "junior"),
    [
        ("Associate Director, Delivery", False),
        ("Associate Product Manager", False),
        ("Associate Lead, Programs", False),
        ("Associate Head of Product", False),
        ("Associate Project Coordinator", True),
        ("Junior Project Coordinator", True),
    ],
)
def test_junior_flag_not_for_senior_titles(title, junior):
    job = make_job("Acme", title)
    assert ("Junior level" in job_flags(job, {}, TODAY)) is junior


# ---------------------------------------------------------------- "safe" is not SAFe
@pytest.mark.parametrize(
    ("red_flags", "flagged"),
    [
        ("Build a safe, inclusive team culture", False),
        ("Safety-critical domain; SAFE environment", False),
        ("SAFe preferred", True),
        ("SAFe Agilist certification required", True),
        ("SAFe POPM a plus", True),
        ("PMP preferred", True),
    ],
)
def test_safe_word_is_not_the_safe_certification(red_flags, flagged):
    job = make_job(**{"Red Flags": red_flags})
    assert ("Certification mentioned" in job_flags(job, {}, TODAY)) is flagged


# ---------------------------------------------------------------- keyword match uses word boundaries
def test_keyword_match_word_boundaries():
    assert not has_word("agile", "fragile systems")
    assert has_word("agile", "agile delivery")
    assert has_word("stakeholder", "worked with stakeholders")
    assert has_word("c++", "python and c++ tools")
    assert keyword_match({"Notes": "Key: agile"}, "maintained fragile legacy code") == 0
    assert keyword_match({"Notes": "Key: agile"}, "ran agile ceremonies") == 100


# ---------------------------------------------------------------- `applied ola` must not mark Motorola
def test_applied_requires_word_prefix_match(tmp_path):
    settings = Settings(base_dir=tmp_path)
    init_tracker(settings)
    add_file(settings, write(tmp_path, [make_job("Motorola", "Project Manager"), make_job("Globex", "PO")]))

    done, problems = mark_applied(settings, ["ola"], date(2026, 9, 29))
    assert done == [] and "0 matches" in problems[0]
    assert all(j["Status"] == "Found" for j in read_jobs(settings.tracker_path))

    done, _ = mark_applied(settings, ["moto"], date(2026, 9, 29))  # prefix of a word is fine
    assert done == ["Motorola - Project Manager"]


def test_applied_ambiguous_changes_nothing(tmp_path):
    settings = Settings(base_dir=tmp_path)
    init_tracker(settings)
    add_file(settings, write(tmp_path, [make_job("Acme", "Project Manager"), make_job("Acme", "Product Owner")]))
    done, problems = mark_applied(settings, ["acme"])
    assert done == [] and "2 matches" in problems[0]
    assert all(j["Status"] == "Found" for j in read_jobs(settings.tracker_path))


# ---------------------------------------------------------------- B-010: no personal data in the package
# Personal terms are never written into the repository, not even in this test. To check for your own
# name, employers, etc., list them (one per line) in tests/.private_terms.txt (git-ignored) or point
# SHORTLIST_PRIVATE_TERMS at such a file.
SKIP_DIRS = {".venv", ".git", "__pycache__", ".pytest_cache", ".ruff_cache", "build", "dist"}
EMAIL_RE = re.compile(r"[\w.+-]+@([\w-]+\.)+[a-z]{2,}", re.I)
ALLOWED_EMAIL_DOMAIN_RE = re.compile(r"@([\w-]+\.)*example\.(com|org|net)$", re.I)
PHONE_RE = re.compile(r"(?<![\w-])\+\d{1,3}[ -]?\d{3,5}[ -]?\d{4,6}(?![\w-])|(?<!\d)\d{10}(?!\d)")


def repo_texts():
    for p in ROOT.rglob("*"):
        if p.is_dir() or SKIP_DIRS & set(p.parts) or p.suffix in (".xlsx", ".pyc") or p.name == ".private_terms.txt":
            continue
        try:
            yield p.relative_to(ROOT), p.read_text(encoding="utf-8")
        except (UnicodeDecodeError, OSError):
            continue


def test_b010_only_example_email_addresses():
    offenders = [(str(p), m.group(0)) for p, text in repo_texts() for m in EMAIL_RE.finditer(text)
                 if not ALLOWED_EMAIL_DOMAIN_RE.search(m.group(0))]
    assert offenders == []


def test_b010_no_real_phone_numbers():
    offenders = [(str(p), m.group(0)) for p, text in repo_texts() for m in PHONE_RE.finditer(text)
                 if set(re.sub(r"\D", "", m.group(0))) != {"0"}]  # all-zero placeholders are fine
    assert offenders == []


def test_b010_private_terms_absent():
    import os

    terms_file = Path(os.environ.get("SHORTLIST_PRIVATE_TERMS", ROOT / "tests" / ".private_terms.txt"))
    if not terms_file.is_file():
        pytest.skip("no private terms file (optional local check)")
    terms = [t.strip().lower() for t in terms_file.read_text(encoding="utf-8").splitlines() if t.strip()]
    offenders = [(str(p), t) for p, text in repo_texts() for t in terms if t in text.lower()]
    assert offenders == []
