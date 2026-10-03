from __future__ import annotations

import json
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from shortlist.tracker import (
    COL,
    TrackerError,
    add_file,
    add_jobs,
    criteria_rows,
    init_tracker,
    load,
    mark_applied,
    new_workbook,
    prune,
    read_jobs,
    rebuild,
    refresh,
    update_file,
)

from .conftest import EXAMPLES, make_job


def iso(days_ago: int) -> str:
    return (date.today() - timedelta(days=days_ago)).isoformat()


def write(tmp_path, name, jobs):
    p = tmp_path / name
    p.write_text(json.dumps(jobs), encoding="utf-8")
    return p


def test_init_creates_tabs_and_refuses_overwrite(settings):
    path = init_tracker(settings)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Dashboard", "Tracker", "Metrics", "Criteria"]
    with pytest.raises(TrackerError):
        init_tracker(settings)


def test_standard_answers_come_from_user_file(settings):
    path = init_tracker(settings, EXAMPLES / "standard_answers.example.json")
    ws = load_workbook(path)["Standard Answers"]
    questions = [r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True)]
    assert "Full name" in questions
    assert "_note" not in questions


def test_criteria_generated_from_settings(settings):
    settings.max_applicants = 300
    settings.criteria = {"Target roles": "Product Owner"}
    text = " ".join(f"{k} {v}" for k, v in criteria_rows(settings))
    assert "300 or fewer" in text
    assert "Target roles Product Owner" in text


def test_add_dedupes_by_url_and_company_title(settings, tmp_path):
    init_tracker(settings)
    jobs = [
        make_job("Acme", "Delivery Manager", **{"Date Posted": iso(1)}),
        make_job("Acme", "Delivery Manager", **{"Job URL": "https://other.example.com/1", "Date Posted": iso(1)}),
        make_job("Acme (for a client)", "Delivery Manager", **{"Job URL": "https://x.example.com/2"}),
        make_job("Globex", "Product Owner", **{"Date Posted": iso(0)}),
    ]
    result = add_file(settings, write(tmp_path, "jobs.json", jobs))
    assert result.added == 2
    assert len(result.duplicates) == 2
    # adding the same file again adds nothing
    assert add_file(settings, write(tmp_path, "jobs.json", jobs)).added == 0
    rows = read_jobs(settings.tracker_path)
    assert {r["Company"] for r in rows} == {"Acme", "Globex"}
    assert all(isinstance(r["Priority"], int) for r in rows)


def test_add_skips_company_in_cooldown(settings):
    wb = new_workbook(settings)
    add_jobs(wb, [make_job("Acme", "Project Manager", Status="Rejected",
                           **{"Response Date": date.today() - timedelta(days=10)})], settings)
    result = add_jobs(wb, [make_job("Acme", "Product Owner"), make_job("Globex", "Product Owner")], settings)
    assert result.added == 1
    assert result.cooldown == ["Acme - Product Owner"]


def test_formula_like_text_is_stored_as_text(settings, tmp_path):
    init_tracker(settings)
    add_file(settings, write(tmp_path, "j.json", [make_job("=HYPERLINK(\"http://evil.example.com\")", "PM")]))
    ws = load_workbook(settings.tracker_path)["Tracker"]
    cell = ws.cell(row=2, column=COL["Company"])
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")


def test_update_by_url(settings, tmp_path):
    init_tracker(settings)
    job = make_job("Acme", "Delivery Manager")
    add_file(settings, write(tmp_path, "a.json", [job]))
    changes = [{"Job URL": job["Job URL"], "Status": "Ready to Apply", "Cover Letter File": "letters/acme.txt"},
               {"Job URL": "https://missing.example.com/x", "Status": "Skipped"}]
    updated, missing = update_file(settings, write(tmp_path, "u.json", changes))
    assert updated == 1 and missing == ["https://missing.example.com/x"]
    row = read_jobs(settings.tracker_path)[0]
    assert row["Status"] == "Ready to Apply"
    assert row["Cover Letter File"] == "letters/acme.txt"


def test_applied_marks_one_match_and_reports_ambiguity(settings, tmp_path):
    init_tracker(settings)
    jobs = [make_job("Acme", "Delivery Manager"), make_job("Acme", "Product Owner"), make_job("Globex", "PM")]
    add_file(settings, write(tmp_path, "a.json", jobs))

    done, problems = mark_applied(settings, ["acme"], date(2026, 9, 29))
    assert done == [] and "2 matches" in problems[0]

    done, problems = mark_applied(settings, ["acme product", "Globex", "nobody"], date(2026, 9, 29))
    assert done == ["Acme - Product Owner", "Globex - PM"]
    assert "'nobody': 0 matches" in problems[0]

    done, _ = mark_applied(settings, [jobs[0]["Job URL"]])
    assert done == ["Acme - Delivery Manager"]

    rows = {r["Role Title"]: r for r in read_jobs(settings.tracker_path)}
    assert rows["Product Owner"]["Status"] == "Applied"
    assert rows["Product Owner"]["Date Applied"] == "2026-09-29"


def test_prune_drops_stale_open_jobs_and_keeps_applied(settings, tmp_path):
    init_tracker(settings, EXAMPLES / "standard_answers.example.json")
    jobs = [
        make_job("Fresh", "PM", **{"Date Posted": iso(1), "Applicants": "Over 200"}),
        make_job("Crowded", "PM", **{"Date Posted": iso(4), "Applicants": "Over 200"}),
        make_job("Quiet", "PM", **{"Date Posted": iso(5), "Applicants": "120"}),
        make_job("Old", "PM", **{"Date Posted": iso(10), "Applicants": "5"}),
        make_job("Skipped", "PM", Status="Skipped", **{"Date Posted": iso(0)}),
        make_job("Applied", "PM", Status="Applied", **{"Date Posted": iso(30), "Date Applied": iso(20)}),
    ]
    add_file(settings, write(tmp_path, "a.json", jobs))
    result = prune(settings)
    assert {j["Company"] for j in result.dropped} == {"Crowded", "Old", "Skipped"}
    assert {j["Company"] for j in read_jobs(settings.tracker_path)} == {"Fresh", "Quiet", "Applied"}
    assert result.backup is not None and result.backup.exists()
    # standard answers survive the rebuild
    assert "Standard Answers" in load(settings).sheetnames


def test_rebuild_and_dashboard(settings, tmp_path):
    init_tracker(settings)
    jobs = [make_job("Acme", "PM", Status="Ready to Apply", **{"Date Posted": iso(0)}),
            make_job("Globex", "PO", Status="Applied", **{"Date Applied": iso(9)})]
    result, backup = rebuild(settings, [write(tmp_path, "r.json", jobs)])
    assert result.added == 2 and backup is not None
    refresh(settings)
    wb = load(settings)
    assert wb.sheetnames[0] == "Dashboard"
    values = [c for row in wb["Dashboard"].iter_rows(values_only=True) for c in row if c is not None]
    assert "Acme" in values  # in "Apply next"
    assert "Globex" in values  # follow-up due (applied 9 days ago)
    assert "Applied in total" in values


def test_missing_tracker_is_a_clear_error(settings):
    with pytest.raises(TrackerError, match="shortlist init"):
        load(settings)


def test_bad_date_is_a_clear_error(settings):
    wb = new_workbook(settings)
    with pytest.raises(TrackerError, match="not a date"):
        add_jobs(wb, [make_job(**{"Date Posted": "last week"})], settings)
