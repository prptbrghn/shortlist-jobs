"""Automated checks for drafted cover letters (``shortlist check-letters``).

Checks are deliberately simple and explainable:
- word count within a range;
- no em or en dashes (a common tell of machine-written text); semicolons are a warning;
- no unfilled placeholders such as ``{COMPANY}``, ``[Company]`` or ``<Role>``;
- none of the quoted phrases listed under a "Never say" heading in the fact sheet;
- none of a short list of cliches (warning only).

These checks do not prove a letter is truthful. A human must still read every letter
against the fact sheet before sending it.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

DEFAULT_MIN_WORDS = 150
DEFAULT_MAX_WORDS = 300
CLICHES = (
    "i am writing to express",
    "hit the ground running",
    "team player",
    "synergy",
    "passionate about",
    "dynamic environment",
    "think outside the box",
    "results-driven",
    "go-getter",
    "i believe i would be a great fit",
)
_PLACEHOLDER_RE = re.compile(
    r"\{[A-Z_]{2,}\}|\[(company|role|name|your name|hiring manager)\]|<(company|role|name)>", re.I
)
_DASH_RE = re.compile("[\u2013\u2014]")


@dataclass
class LetterReport:
    path: Path
    words: int
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def never_say_phrases(fact_sheet: str) -> list[str]:
    """Quoted phrases under a "Never say" heading in the fact sheet (case-insensitive heading match)."""
    match = re.search(r"^#+\s*never say\s*$(.*?)(?=^#+\s|\Z)", fact_sheet, re.I | re.M | re.S)
    if not match:
        return []
    return [p.strip() for p in re.findall(r"[\"\u201c]([^\"\u201d]+)[\"\u201d]", match.group(1)) if p.strip()]


def check_letter_text(text: str, never_say: list[str] | None = None, min_words: int = DEFAULT_MIN_WORDS,
                      max_words: int = DEFAULT_MAX_WORDS, path: Path | None = None) -> LetterReport:
    """Check one letter's text and return a report."""
    words = len(re.findall(r"\b[\w'-]+\b", text))
    report = LetterReport(path=path or Path("<text>"), words=words)
    if not min_words <= words <= max_words:
        report.errors.append(f"{words} words (expected {min_words}-{max_words})")
    if _DASH_RE.search(text):
        report.errors.append("contains an em or en dash")
    if ";" in text:
        report.warnings.append("contains a semicolon")
    placeholder = _PLACEHOLDER_RE.search(text)
    if placeholder:
        report.errors.append(f"unfilled placeholder {placeholder.group(0)!r}")
    lowered = text.lower()
    for phrase in never_say or []:
        if phrase.lower() in lowered:
            report.errors.append(f"uses a 'Never say' phrase: {phrase!r}")
    for cliche in CLICHES:
        if cliche in lowered:
            report.warnings.append(f"cliche: {cliche!r}")
    return report


def check_letters(paths: list[Path], fact_sheet: Path | None = None, min_words: int = DEFAULT_MIN_WORDS,
                  max_words: int = DEFAULT_MAX_WORDS) -> list[LetterReport]:
    """Check every letter file; ``.txt`` and ``.md`` are read directly, ``.docx`` via python-docx."""
    never_say = never_say_phrases(fact_sheet.read_text(encoding="utf-8-sig")) if fact_sheet else []
    reports = []
    for p in paths:
        if p.suffix.lower() == ".docx":
            import docx  # python-docx

            text = "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
        else:
            text = p.read_text(encoding="utf-8-sig")
        reports.append(check_letter_text(text, never_say, min_words, max_words, p))
    return reports
