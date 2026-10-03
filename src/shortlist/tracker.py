"""Excel job tracker: create, add, update, prune, mark applied, rebuild, and refresh the dashboard.

The workbook has these tabs:
    Dashboard         static summary, rebuilt on every save (apply next, expiring, follow-ups, top N)
    Tracker           one row per job; ID, Days Old and Follow-up Date are Excel formulas
    Metrics           weekly and per-source / resume / role response rates (Excel formulas)
    Criteria          your rules, generated from settings.json
    Standard Answers  optional; copied from your own answers file (never shipped with the package)

Every destructive operation (prune, rebuild) moves the old file into ``backups/`` first.
"""

from __future__ import annotations

import json
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .config import Settings
from .rules import as_date, company_key, is_stale, job_key
from .safety import is_safe_url
from .scoring import SERIOUS_FLAGS, job_flags, keyword_match, priority, tailored_resume

ROWS = 1000
DATE_COLUMNS = ("Date Found", "Date Posted", "Date Applied", "Response Date")
FORMULA_COLUMNS = ("ID", "Days Old", "Follow-up Date")

COLUMNS: list[tuple[str, int]] = [
    ("ID", 6), ("Date Found", 12), ("Date Posted", 12), ("Days Old", 8), ("Applicants", 14),
    ("Company", 22), ("Role Title", 32), ("Role Type", 17),
    ("Source", 14), ("Job URL", 32), ("Location", 20), ("Work Mode", 11), ("Country", 13),
    ("Visa / Sponsorship", 20), ("Salary Listed", 18), ("Experience Required", 13),
    ("Match Score (1-10)", 10), ("Priority", 9), ("Keyword Match", 10), ("Flags", 34), ("Why It Fits", 45),
    ("Red Flags", 32), ("Resume Version", 17),
    ("Cover Letter File", 30), ("Status", 14), ("Date Applied", 12), ("Follow-up Date", 12),
    ("Response Date", 12), ("Response Type", 16), ("Contact / Referral", 22), ("Notes", 40),
]
COL = {name: i + 1 for i, (name, _) in enumerate(COLUMNS)}
LETTER = {name: get_column_letter(i) for name, i in COL.items()}

DEFAULT_LISTS: dict[str, list[str]] = {
    "Role Type": ["Project Manager", "Delivery Manager", "Product Manager", "Product Owner"],
    "Source": ["LinkedIn", "Indeed", "Wellfound", "Glassdoor", "Remote boards", "Company site", "Referral", "Other"],
    "Work Mode": ["Remote", "Hybrid", "Onsite"],
    "Visa / Sponsorship": ["Not needed", "Remote from home country", "Sponsorship offered", "Sponsorship unclear"],
    "Status": ["Found", "Shortlisted", "Ready to Apply", "Applied", "Screening", "Interviewing", "Offer",
               "Rejected", "Withdrawn", "Skipped"],
    "Response Type": ["None", "Recruiter Call", "Assessment", "Interview", "Rejection", "Offer"],
}
OPEN_STATUSES = (None, "", "Found", "Shortlisted", "Ready to Apply", "Skipped")

NAVY = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
DATE_FMT = "dd-mmm-yyyy"


class TrackerError(RuntimeError):
    """Raised for tracker problems the user can fix (missing file, full sheet, bad input)."""


# ---------------------------------------------------------------- helpers


def lists_for(settings: Settings) -> dict[str, list[str]]:
    """Dropdown lists: defaults, overridden by settings.lists; resume versions come from settings.resumes."""
    lists = {**DEFAULT_LISTS, **settings.lists}
    lists["Resume Version"] = [*settings.resumes, "Tailored"]
    return lists


FORMULA_PREFIXES = ("=", "+", "-", "@")


def put(ws: Worksheet, row: int, column: int, value: Any) -> Any:
    """Write a value safely (B-004: spreadsheet formula injection).

    Text starting with =, +, - or @ is stored as a text cell with Excel's quote prefix, so it is
    never evaluated as a formula, even if the user later edits the cell.
    """
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        cell.data_type = "s"  # openpyxl would otherwise store a leading "=" as a formula
        cell.quotePrefix = True
    return cell


def header(ws: Worksheet, names: list[str], row: int = 1) -> None:
    for i, name in enumerate(names, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill = HEAD_FONT, NAVY
        c.alignment = Alignment(wrap_text=True, vertical="center")


def read_json_list(path: str | Path) -> list[dict[str, Any]]:
    data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if not isinstance(data, list) or not all(isinstance(j, dict) for j in data):
        raise TrackerError(f"{path}: expected a JSON array of objects")
    return data


# ---------------------------------------------------------------- building sheets


def build_tracker(ws: Worksheet, settings: Settings) -> None:
    ws.title = "Tracker"
    header(ws, [n for n, _ in COLUMNS])
    ws.row_dimensions[1].height = 32
    for name, width in COLUMNS:
        ws.column_dimensions[LETTER[name]].width = width
    ws.freeze_panes = f"{LETTER['Role Type']}2"
    ws.auto_filter.ref = f"A1:{LETTER['Notes']}{ROWS}"

    for name, items in lists_for(settings).items():
        dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{LETTER[name]}2:{LETTER[name]}{ROWS}")
    score = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True)
    ws.add_data_validation(score)
    score.add(f"{LETTER['Match Score (1-10)']}2:{LETTER['Match Score (1-10)']}{ROWS}")

    a, t, u = LETTER["Date Applied"], LETTER["Follow-up Date"], LETTER["Response Type"]
    s, company = LETTER["Status"], LETTER["Company"]
    posted, age = LETTER["Date Posted"], LETTER["Days Old"]
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ROWS + 1):
        ws[f"A{r}"] = f'=IF({company}{r}="","",ROW()-1)'
        ws[f"{age}{r}"] = f'=IF({posted}{r}="","",TODAY()-{posted}{r})'
        ws[f"{age}{r}"].number_format = "0"
        ws[f"{t}{r}"] = f'=IF({a}{r}="","",{a}{r}+{settings.follow_up_days})'
        for col in (*DATE_COLUMNS, "Follow-up Date"):
            ws[f"{LETTER[col]}{r}"].number_format = DATE_FMT
        ws[f"{LETTER['Job URL']}{r}"].font = LINK_FONT
        for col in ("Why It Fits", "Red Flags", "Notes", "Flags"):
            ws[f"{LETTER[col]}{r}"].alignment = wrap

    last = LETTER["Notes"]
    fresh, window = settings.fresh_days, settings.window_days
    # posting age: green = fresh, amber = inside the window, red = older than the window
    for formula, color in ((f'AND(${age}2<>"",${age}2<={fresh})', "C6EFCE"),
                           (f'AND(${age}2<>"",${age}2>{fresh},${age}2<={window})', "FFEB9C"),
                           (f'AND(${age}2<>"",${age}2>{window})', "FFC7CE")):
        ws.conditional_formatting.add(f"{age}2:{age}{ROWS}",
                                      FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add(
        f"A2:{last}{ROWS}",
        FormulaRule(formula=[f'AND(${t}2<>"",${t}2<=TODAY(),${u}2="")'], fill=PatternFill("solid", fgColor="FFF2CC")))
    ws.conditional_formatting.add(
        f"{s}2:{s}{ROWS}",
        FormulaRule(formula=[f'OR(${s}2="Screening",${s}2="Interviewing",${s}2="Offer")'],
                    fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(
        f"A2:{last}{ROWS}",
        FormulaRule(formula=[f'OR(${s}2="Rejected",${s}2="Skipped",${s}2="Withdrawn")'], font=Font(color="808080")))


def build_metrics(ws: Worksheet, settings: Settings, start: date | None = None) -> None:
    ws.title = "Metrics"
    a_col, w_col, e = LETTER["Date Applied"], LETTER["Response Type"], "Tracker!"
    applied = f"{e}${a_col}:${a_col}"
    resp = f"{e}${w_col}:${w_col}"
    responded = f'{resp},"<>",{resp},"<>None"'

    ws["A1"] = "Weekly"
    ws["A1"].font = Font(bold=True, size=13)
    header(ws, ["Week Starting", "Applied", "Responses", "Response Rate", "Interviews"], row=2)
    today = start or date.today()
    monday = today - timedelta(days=today.weekday())
    for i in range(16):
        r = 3 + i
        ws[f"A{r}"] = monday + timedelta(weeks=i)
        ws[f"A{r}"].number_format = DATE_FMT
        rng = f'{applied},">="&A{r},{applied},"<"&A{r}+7'
        ws[f"B{r}"] = f"=COUNTIFS({rng})"
        ws[f"C{r}"] = f"=COUNTIFS({rng},{responded})"
        ws[f"D{r}"] = f'=IF(B{r}=0,"",C{r}/B{r})'
        ws[f"D{r}"].number_format = "0%"
        ws[f"E{r}"] = f'=COUNTIFS({rng},{resp},"Interview")'

    lists = lists_for(settings)
    top = 21
    for title, col in (("By Source", "Source"), ("By Resume Version", "Resume Version"), ("By Role Type", "Role Type")):
        ws[f"A{top}"] = title
        ws[f"A{top}"].font = Font(bold=True, size=13)
        header(ws, [title, "Applied", "Responses", "Response Rate", "Interviews"], row=top + 1)
        key = f"{e}${LETTER[col]}:${LETTER[col]}"
        first = top + 2
        items = lists[col]
        for i, item in enumerate(items):
            r = first + i
            ws[f"A{r}"] = item
            ws[f"B{r}"] = f'=COUNTIFS({key},A{r},{applied},"<>")'
            ws[f"C{r}"] = f'=COUNTIFS({key},A{r},{applied},"<>",{responded})'
            ws[f"D{r}"] = f'=IF(B{r}=0,"",C{r}/B{r})'
            ws[f"D{r}"].number_format = "0%"
            ws[f"E{r}"] = f'=COUNTIFS({key},A{r},{resp},"Interview")'
        total = first + len(items)
        ws[f"A{total}"] = "Total"
        for c in "BCE":
            ws[f"{c}{total}"] = f"=SUM({c}{first}:{c}{total - 1})"
        ws[f"D{total}"] = f'=IF(B{total}=0,"",C{total}/B{total})'
        ws[f"D{total}"].number_format = "0%"
        for c in "ABCDE":
            ws[f"{c}{total}"].font = Font(bold=True)
        top = total + 3

    ws.column_dimensions["A"].width = 22
    for c in "BCDE":
        ws.column_dimensions[c].width = 14


def build_table(ws: Worksheet, title: str, rows: list[tuple[str, str]], widths: tuple[int, int] = (38, 100)) -> None:
    ws.title = title
    header(ws, list(rows[0]))
    for r, (k, v) in enumerate(rows[1:], start=2):
        put(ws, r, 1, k).font = Font(bold=True)
        put(ws, r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = widths


def criteria_rows(settings: Settings) -> list[tuple[str, str]]:
    """The Criteria tab: the user's own criteria text plus the rules, generated from settings."""
    s, w = settings, settings.priority_weights
    fresh, window = s.fresh_days, s.window_days
    rows: list[tuple[str, str]] = [("Rule", "Value")]
    rows += [(k, v) for k, v in s.criteria.items()]
    by_role = "; ".join(f"{role} -> {version}" for role, version in s.resume_by_role_type.items())
    rows += [
        ("Freshness + applicants",
         f"Ages are calendar days, not hours. Posted today or within the previous {fresh} calendar day(s): keep "
         f"regardless of applicants. Posted {fresh + 1}-{window} days ago: keep ONLY if a specific applicant count "
         f"of {s.max_applicants} or fewer is shown ('Over 200', '200+' or no count = drop). Older than {window} days, "
         "undated or dated in the future: drop. Use the ATS first-published date to catch reposts. "
         f"'Days Old' column: green <={fresh}, amber {fresh + 1}-{window}, red >{window}"),
        ("Priority score (0-100)",
         f"Fit (Match Score x {w['fit']:g}, max {10 * w['fit']:g}) + freshness (max {w['freshness']:g}, falls to 0 at "
         f"{window} days) + competition (max {w['competition']:g}: <=50 applicants 100%, <=100 75%, <=200 50%, "
         f"unknown 25%, more 10%) + resume keyword match (max {w['keywords']:g}) - {w['serious_flag_penalty']:g} per "
         f"serious flag ({', '.join(SERIOUS_FLAGS)}). Weights live in settings.json"),
        ("Flags",
         "Automatic warnings: recruiter or hidden employer, possible repost, US or night hours, junior level, "
         f"asks {s.flag_experience_years}+ years, certification mentioned, visa unclear, check pay vs floor, "
         "already applied here, rejected here recently, suspicious page text"),
        ("Keyword Match",
         "Share of the job's key requirements whose main words appear in the resume version you will send"),
        ("Same company", f"Skip companies that rejected you in the last {s.reapply_cooldown_days} days; "
                         "flag if you already applied to another role there"),
        ("Follow-up", f"Follow up {s.follow_up_days} days after applying if there is no response"),
        ("Daily shortlist", f"Top {s.top_n} open jobs by priority"),
        ("Resume versions", ", ".join(s.resumes) + (f" (by role type: {by_role})" if by_role else "")),
        ("Human in the loop", "Shortlist drafts and ranks. You review, edit and submit every application yourself"),
    ]
    return rows


def load_answers(path: str | Path) -> list[tuple[str, str]]:
    """Read Standard Answers from JSON: an object {question: answer}, or a list of pairs / objects."""
    data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if isinstance(data, dict):
        pairs = list(data.items())
    elif isinstance(data, list):
        pairs = []
        for item in data:
            if isinstance(item, dict) and "Question" in item:
                pairs.append((item["Question"], item.get("Answer", "")))
            elif isinstance(item, (list, tuple)) and len(item) == 2:
                pairs.append((item[0], item[1]))
            else:
                raise TrackerError(f"{path}: each answer must be [question, answer] or {{'Question', 'Answer'}}")
    else:
        raise TrackerError(f"{path}: expected a JSON object or array")
    # Keys starting with "_" are comments.
    return [(str(q), str(a)) for q, a in pairs if str(q).strip() and not str(q).startswith("_")]


def _answers_from_workbook(path: Path) -> list[tuple[str, str]] | None:
    try:
        wb = load_workbook(path)
    except Exception:  # an unreadable old file just means "no answers to carry over"
        return None
    if "Standard Answers" not in wb.sheetnames:
        return None
    ws = wb["Standard Answers"]
    return [(str(r[0]), "" if r[1] is None else str(r[1]))
            for r in ws.iter_rows(min_row=2, max_col=2, values_only=True) if r[0]]


def new_workbook(settings: Settings, answers: list[tuple[str, str]] | None = None) -> Workbook:
    """Build an empty tracker workbook (not saved)."""
    wb = Workbook()
    build_tracker(wb.active, settings)
    build_metrics(wb.create_sheet(), settings)
    build_table(wb.create_sheet(), "Criteria", criteria_rows(settings))
    if answers:
        build_table(wb.create_sheet(), "Standard Answers", [("Question", "Answer"), *answers], widths=(34, 100))
    return wb


# ---------------------------------------------------------------- reading and writing rows


def write_fields(ws: Worksheet, row: int, job: dict[str, Any]) -> None:
    for name, value in job.items():
        if name not in COL or name in FORMULA_COLUMNS or value in ("", None):
            continue
        if name in DATE_COLUMNS and not isinstance(value, (date, datetime)):
            parsed = as_date(value)
            if parsed is None:
                raise TrackerError(f"{job.get('Company')} - {job.get('Role Title')}: '{name}' is not a date: {value!r}")
            value = parsed
        cell = put(ws, row, COL[name], value)
        if name == "Job URL" and is_safe_url(value):
            cell.hyperlink = value


def url_rows(ws: Worksheet) -> dict[Any, int]:
    col = COL["Job URL"]
    return {ws.cell(row=r, column=col).value: r for r in range(2, ROWS + 1) if ws.cell(row=r, column=col).value}


def row_jobs(ws: Worksheet) -> list[tuple[int, dict[str, Any]]]:
    """(row number, {column: value}) for every filled row."""
    names = [c.value for c in ws[1]]
    out = []
    for r in range(2, ROWS + 1):
        if not ws.cell(row=r, column=COL["Company"]).value:
            continue
        out.append((r, {n: ws.cell(row=r, column=i + 1).value for i, n in enumerate(names) if n}))
    return out


def read_jobs(path: Path) -> list[dict[str, Any]]:
    """All jobs in a tracker file, with dates as ISO strings and formula columns removed."""
    jobs = []
    for _, job in row_jobs(load_workbook(path)["Tracker"]):
        for n in DATE_COLUMNS:
            if isinstance(job.get(n), datetime):
                job[n] = job[n].date().isoformat()
        for n in FORMULA_COLUMNS:
            job.pop(n, None)
        jobs.append(job)
    return jobs


@dataclass
class AddResult:
    added: int = 0
    duplicates: list[str] = field(default_factory=list)
    cooldown: list[str] = field(default_factory=list)
    invalid: list[str] = field(default_factory=list)


def add_jobs(wb: Workbook, jobs: Iterable[dict[str, Any]], settings: Settings, today: date | None = None) -> AddResult:
    """Append jobs, skipping duplicate links, duplicate company + title, and companies in the reapply cooldown."""
    ws = wb["Tracker"]
    today = today or date.today()
    seen = url_rows(ws)
    existing = row_jobs(ws)
    keys = {job_key(j.get("Company"), j.get("Role Title")) for _, j in existing}
    cutoff = today - timedelta(days=settings.reapply_cooldown_days)
    recently_rejected = {
        company_key(j.get("Company")) for _, j in existing
        if j.get("Status") == "Rejected"
        and (as_date(j.get("Response Date")) or as_date(j.get("Date Applied")) or today) >= cutoff
    }
    rows = (r for r in range(2, ROWS + 1) if ws.cell(row=r, column=COL["Company"]).value in (None, ""))
    result = AddResult()
    for job in jobs:
        label = f"{job.get('Company')} - {job.get('Role Title')}"
        if not job.get("Company"):
            result.invalid.append(f"{label} (no company)")
            continue
        key = job_key(job.get("Company"), job.get("Role Title"))
        if (job.get("Job URL") and job.get("Job URL") in seen) or key in keys:
            result.duplicates.append(label)
            continue
        if company_key(job.get("Company")) in recently_rejected:
            result.cooldown.append(label)
            continue
        row = next(rows, None)
        if row is None:
            raise TrackerError(f"the tracker is full ({ROWS - 1} rows); run 'shortlist prune' first")
        write_fields(ws, row, job)
        if job.get("Job URL"):
            seen[job.get("Job URL")] = row
        keys.add(key)
        result.added += 1
    return result


# ---------------------------------------------------------------- computed columns and dashboard


class ResumeTexts:
    """Loads resume .docx text once per resume version for keyword matching."""

    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self._cache: dict[str, str] = {}

    def path_for(self, name: str) -> Path:
        p = Path(name)
        if p.suffix.lower() != ".docx":
            p = p.with_name(p.name + ".docx")
        return p if p.is_absolute() else self.settings.resume_dir / p

    def text(self, version: str | None, notes: Any = None) -> str:
        s = self.settings
        name = tailored_resume(notes) or s.resumes.get(version or "") or s.resumes[s.default_resume_version]
        if name not in self._cache:
            path = self.path_for(name)
            try:
                import docx  # python-docx

                self._cache[name] = (" ".join(p.text for p in docx.Document(str(path)).paragraphs).lower()
                                     if path.exists() else "")
            except Exception:  # unreadable resume: keyword match is simply left blank
                self._cache[name] = ""
        return self._cache[name]


def refresh_computed(wb: Workbook, settings: Settings, today: date | None = None,
                     resumes: ResumeTexts | None = None) -> None:
    """Recompute Flags, Keyword Match and Priority for every row."""
    ws = wb["Tracker"]
    today = today or date.today()
    resumes = resumes or ResumeTexts(settings)
    rows = row_jobs(ws)
    by_company: dict[str, list[dict[str, Any]]] = {}
    for _, job in rows:
        by_company.setdefault(company_key(job.get("Company")), []).append(job)
    for r, job in rows:
        flags = job_flags(job, by_company, today, settings)
        km = keyword_match(job, resumes.text(job.get("Resume Version"), job.get("Notes")))
        put(ws, r, COL["Flags"], ", ".join(flags) or None)
        ws.cell(row=r, column=COL["Keyword Match"], value=km / 100 if km is not None else None).number_format = "0%"
        ws.cell(row=r, column=COL["Priority"], value=priority(job, flags, km, today, settings))


def build_dashboard(wb: Workbook, settings: Settings, today: date | None = None) -> None:
    """A static summary sheet, rebuilt every time the tracker is saved by Shortlist."""
    today = today or date.today()
    jobs = []
    for r, vals in row_jobs(wb["Tracker"]):
        vals["_row"] = r
        jobs.append(vals)
    window, fu_days = settings.window_days, settings.follow_up_days
    for j in jobs:
        posted = as_date(j.get("Date Posted"))
        j["_left"] = (window - (today - posted).days) if posted else None
        applied_on = as_date(j.get("Date Applied"))
        j["_follow"] = applied_on + timedelta(days=fu_days) if applied_on else None
    open_ = [j for j in jobs if j.get("Status") in OPEN_STATUSES]
    applied = [j for j in jobs if j.get("Status") not in OPEN_STATUSES]
    responded = [j for j in applied if j.get("Response Type") not in (None, "", "None")]
    week_start = today - timedelta(days=today.weekday())
    this_week = [j for j in applied if (as_date(j.get("Date Applied")) or date.min) >= week_start]

    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]
    d = wb.create_sheet("Dashboard", 0)
    wb.active = 0
    for col, width in zip("ABCDEFGH", (26, 38, 10, 10, 16, 34, 40, 46), strict=True):
        d.column_dimensions[col].width = width
    d["A1"] = "Job Search Dashboard"
    d["A1"].font = Font(bold=True, size=16, color="1F3864")
    d["A2"] = f"Updated {today:%d %b %Y}. Refreshed automatically whenever Shortlist updates the tracker."
    d["A2"].font = Font(italic=True, color="666666")

    stats = [("Ready to apply", sum(j.get("Status") == "Ready to Apply" for j in jobs)),
             ("Shortlisted", sum(j.get("Status") == "Shortlisted" for j in jobs)),
             ("Other open jobs", sum(j.get("Status") in (None, "", "Found") for j in jobs)),
             ("Applied this week", len(this_week)),
             ("Applied in total", len(applied)),
             ("Responses", len(responded)),
             ("Response rate", f"{len(responded) / len(applied):.0%}" if applied else "-")]
    for i, (k, v) in enumerate(stats, start=4):
        d.cell(row=i, column=1, value=k).font = Font(bold=True)
        d.cell(row=i, column=2, value=v)

    Col = tuple[str, Callable[[dict[str, Any]], Any]]

    def table(top: int, title: str, rows: list[dict[str, Any]], cols: list[Col]) -> int:
        d.cell(row=top, column=1, value=title).font = Font(bold=True, size=13, color="1F3864")
        if not rows:
            d.cell(row=top + 1, column=1, value="Nothing here right now.").font = Font(italic=True, color="666666")
            return top + 3
        header(d, [c[0] for c in cols], row=top + 1)
        for i, j in enumerate(rows, start=top + 2):
            for c, (_, fn) in enumerate(cols, start=1):
                v = fn(j)
                cell = put(d, i, c, v)
                if isinstance(v, str) and is_safe_url(v):
                    cell.hyperlink = v
                    cell.font = LINK_FONT
        return top + len(rows) + 3

    def order(j: dict[str, Any]) -> tuple[int, int]:
        return (j["_left"] if j["_left"] is not None else 99, -(j.get("Priority") or 0))

    base: list[Col] = [("Company", lambda j: j.get("Company")), ("Role", lambda j: j.get("Role Title"))]
    link: Col = ("Job link", lambda j: j.get("Job URL"))

    top = 13
    top = table(top, "1. Apply next (Ready to Apply, soonest to expire first)",
                sorted([j for j in jobs if j.get("Status") == "Ready to Apply"], key=order),
                [*base, ("Priority", lambda j: j.get("Priority")), ("Days left", lambda j: j["_left"]),
                 ("Applicants", lambda j: j.get("Applicants")), ("Flags", lambda j: j.get("Flags")),
                 link, ("Cover letter", lambda j: j.get("Cover Letter File"))])
    top = table(top, "2. Expiring within a day (not applied yet)",
                sorted([j for j in open_ if j["_left"] is not None and j["_left"] <= 1
                        and j.get("Status") not in ("Ready to Apply", "Skipped")], key=order),
                [*base, ("Posted", lambda j: as_date(j.get("Date Posted"))), ("Days left", lambda j: j["_left"]),
                 ("Status", lambda j: j.get("Status")), link])
    top = table(top, f"3. Follow-ups due ({fu_days} days after applying, no response yet)",
                sorted([j for j in applied if j["_follow"] and j["_follow"] <= today
                        and j.get("Response Type") in (None, "", "None") and j.get("Status") == "Applied"],
                       key=lambda j: j["_follow"]),
                [*base, ("Applied", lambda j: as_date(j.get("Date Applied"))),
                 ("Follow up since", lambda j: j["_follow"]),
                 ("Contact", lambda j: j.get("Contact / Referral")), link])
    table(top, f"4. Top {settings.top_n} open jobs by priority",
          sorted([j for j in open_ if j.get("Status") != "Skipped"],
                 key=lambda j: (-(j.get("Priority") or 0), order(j)))[: settings.top_n],
          [*base, ("Priority", lambda j: j.get("Priority")), ("Days left", lambda j: j["_left"]),
           ("Status", lambda j: j.get("Status")), ("Flags", lambda j: j.get("Flags")), link])
    for row in d.iter_rows():
        for c in row:
            if isinstance(c.value, date):
                c.number_format = DATE_FMT


# ---------------------------------------------------------------- file operations


def _path(settings: Settings, path: Path | None) -> Path:
    return Path(path) if path else settings.tracker_path


def load(settings: Settings, path: Path | None = None) -> Workbook:
    p = _path(settings, path)
    if not p.exists():
        raise TrackerError(f"{p} does not exist; run 'shortlist init' first")
    return load_workbook(p)


def save(wb: Workbook, settings: Settings, path: Path | None = None, today: date | None = None) -> Path:
    """Recompute Flags / Keyword Match / Priority, rebuild the Dashboard and write the file."""
    p = _path(settings, path)
    refresh_computed(wb, settings, today)
    build_dashboard(wb, settings, today)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def init_tracker(settings: Settings, answers_path: str | Path | None = None, path: Path | None = None) -> Path:
    """Create a new tracker. Refuses to overwrite an existing file."""
    p = _path(settings, path)
    if p.exists():
        raise TrackerError(f"{p} already exists; not overwriting")
    answers_file = answers_path or settings.standard_answers_path
    answers = load_answers(answers_file) if answers_file else None
    return save(new_workbook(settings, answers), settings, p)


def back_up(settings: Settings, path: Path | None = None) -> Path:
    """Move the tracker into backups/ with a timestamp. Returns the backup path."""
    p = _path(settings, path)
    folder = p.parent / "backups"
    folder.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup = folder / f"{p.stem}_{stamp}{p.suffix}"
    p.rename(backup)
    return backup


def _recreate(settings: Settings, jobs: list[dict[str, Any]], path: Path | None,
              today: date | None) -> tuple[AddResult, Path | None]:
    """Back up the current tracker (if any) and write a fresh one containing ``jobs``."""
    p = _path(settings, path)
    answers = None
    backup = None
    if p.exists():
        answers = _answers_from_workbook(p)
        backup = back_up(settings, p)
    if answers is None and settings.standard_answers_path:
        answers = load_answers(settings.standard_answers_path)
    wb = new_workbook(settings, answers)
    result = add_jobs(wb, jobs, settings, today)
    save(wb, settings, p, today)
    return result, backup


def add_file(settings: Settings, json_path: str | Path, path: Path | None = None,
             today: date | None = None) -> AddResult:
    """``shortlist add``: append jobs from a JSON array."""
    jobs = read_json_list(json_path)
    wb = load(settings, path)
    result = add_jobs(wb, jobs, settings, today)
    save(wb, settings, path, today)
    return result


def update_file(settings: Settings, json_path: str | Path, path: Path | None = None) -> tuple[int, list[Any]]:
    """``shortlist update``: set fields on existing rows matched by "Job URL". Returns (updated, missing URLs)."""
    changes = read_json_list(json_path)
    wb = load(settings, path)
    ws = wb["Tracker"]
    rows = url_rows(ws)
    missing = [c.get("Job URL") for c in changes if c.get("Job URL") not in rows]
    for c in changes:
        if c.get("Job URL") in rows:
            write_fields(ws, rows[c["Job URL"]], c)
    save(wb, settings, path)
    return len(changes) - len(missing), missing


def find_rows(ws: Worksheet, name: str) -> list[int]:
    """Rows matching ``name`` (case-insensitive): the exact job link, or text that starts at a word
    boundary in "company role title" ("acme", "acme product", "blue" for "Bluebird").

    Substring-only matches do not count: "ola" does not match "Motorola".
    """
    n = " ".join(name.lower().split())
    if not n:
        return []
    word_prefix = re.compile(rf"(?<![a-z0-9]){re.escape(n)}")
    hits = []
    for r in range(2, ROWS + 1):
        company = ws.cell(row=r, column=COL["Company"]).value
        if not company:
            continue
        title = ws.cell(row=r, column=COL["Role Title"]).value or ""
        url = str(ws.cell(row=r, column=COL["Job URL"]).value or "").lower()
        text = " ".join(f"{company} {title}".lower().split())
        if n == url or word_prefix.search(text):
            hits.append(r)
    return hits


def mark_applied(settings: Settings, names: Iterable[str], when: date | None = None,
                 path: Path | None = None) -> tuple[list[str], list[str]]:
    """Mark jobs Applied (setting Date Applied if empty). Returns (done, problems).

    Each name must match exactly one row (company, part of "company role", or the job link);
    zero or several matches are reported and nothing is changed for that name.
    """
    when = when or date.today()
    wb = load(settings, path)
    ws = wb["Tracker"]
    done, problems = [], []
    for name in names:
        hits = find_rows(ws, name)
        if len(hits) != 1:
            found = [f"{ws.cell(row=r, column=COL['Company']).value} - {ws.cell(row=r, column=COL['Role Title']).value}"
                     for r in hits]
            problems.append(f"'{name}': {len(hits)} matches {found}")
            continue
        r = hits[0]
        ws.cell(row=r, column=COL["Status"], value="Applied")
        if not ws.cell(row=r, column=COL["Date Applied"]).value:
            ws.cell(row=r, column=COL["Date Applied"], value=when)
        done.append(f"{ws.cell(row=r, column=COL['Company']).value} - {ws.cell(row=r, column=COL['Role Title']).value}")
    save(wb, settings, path)
    return done, problems


def refresh(settings: Settings, path: Path | None = None) -> Path:
    """``shortlist dashboard``: recompute scores and rebuild the Dashboard tab."""
    return save(load(settings, path), settings, path)


@dataclass
class PruneResult:
    kept: list[dict[str, Any]]
    dropped: list[dict[str, Any]]
    backup: Path | None


def prune(settings: Settings, path: Path | None = None, today: date | None = None) -> PruneResult:
    """Drop jobs not yet acted on that fail the freshness rule or were Skipped (backs up first).

    Anything applied to or in progress is always kept.
    """
    p = _path(settings, path)
    if not p.exists():
        raise TrackerError(f"{p} does not exist; run 'shortlist init' first")
    today = today or date.today()
    keep, dropped = [], []
    for job in read_jobs(p):
        stale = is_stale(job, today, settings)
        (dropped if job.get("Status") in OPEN_STATUSES and stale else keep).append(job)
    _, backup = _recreate(settings, keep, p, today)
    return PruneResult(keep, dropped, backup)


def rebuild(settings: Settings, json_paths: Iterable[str | Path], path: Path | None = None,
            today: date | None = None) -> tuple[AddResult, Path | None]:
    """Start a fresh tracker containing only the jobs in these JSON files (backs up the old file)."""
    jobs = [job for jp in json_paths for job in read_json_list(jp)]
    return _recreate(settings, jobs, path, today)
